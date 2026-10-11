"""
PDQE Stage 02b — Layer 2 content sanitisation tests.

Groups:
  1. Size limit
  2. CSV sanitisation
  3. JSON sanitisation
  4. TXT sanitisation
  5. XML sanitisation
  6. XLSX sanitisation
  7. DOCX sanitisation
  8. PDF sanitisation
  9. content_hash
 10. Unknown extension and Layer2Result contract

Every fixture file is built in memory — no binary fixtures are committed.
"""

from __future__ import annotations

import hashlib
import io
import json
import zipfile

import docx
import openpyxl
import pytest
from pypdf import PdfReader, PdfWriter

from app.security import layer2_sanitisation as layer2
from app.security.layer2_sanitisation import MAX_BYTES, Layer2Result, run_layer2


# ── Fixture builders ──────────────────────────────────────────────────────────

def _xlsx(rows: list[list], title: str = "Catch") -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = title
    for row in rows:
        sheet.append(row)
    buf = io.BytesIO()
    workbook.save(buf)
    return buf.getvalue()


def _docx(paragraphs: list[str]) -> bytes:
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _pdf(page_texts: list[str]) -> bytes:
    """
    Minimal valid PDF with one Helvetica text line per page and a JavaScript
    OpenAction in the catalog (which Layer 2 must never carry through).
    """
    n = len(page_texts)
    font_id = 3 + 2 * n
    js_id = font_id + 1
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(n))
    objects = [
        f"<< /Type /Catalog /Pages 2 0 R /OpenAction {js_id} 0 R >>".encode(),
        f"<< /Type /Pages /Kids [{kids}] /Count {n} >>".encode(),
    ]
    for i, text in enumerate(page_texts):
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {4 + 2 * i} 0 R "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode()
        )
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    objects.append(b"<< /S /JavaScript /JS (app.alert\\('pwned'\\)) >>")

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    return bytes(out)


def _text(result: Layer2Result) -> str:
    assert result.passed, result.rejection_reason
    assert result.sanitised_content is not None
    return result.sanitised_content.decode("utf-8")


# ── 1. Size limit ─────────────────────────────────────────────────────────────

def test_oversized_content_rejected():
    result = run_layer2(b"a" * (MAX_BYTES + 1), "txt")
    assert not result.passed
    assert result.content_hash is None and result.sanitised_content is None
    assert result.rejection_reason.startswith("content exceeds 50MB limit")
    assert str(MAX_BYTES + 1) in result.rejection_reason


def test_max_size_boundary_accepted():
    result = run_layer2(b"a" * MAX_BYTES, "txt")
    assert result.passed
    assert len(result.sanitised_content) == MAX_BYTES


# ── 2. CSV ────────────────────────────────────────────────────────────────────

def test_csv_clean_content_passes():
    raw = b"year,country,catch_t\r\n2024,WSM,-3.5\r\n2024,FJI,+12\r\n"
    text = _text(run_layer2(raw, "csv"))
    # Normalised to LF; signed numbers are data, not formulae, and survive.
    assert text == "year,country,catch_t\n2024,WSM,-3.5\n2024,FJI,+12\n"


def test_csv_formula_injection_stripped():
    raw = b"name,value\n=cmd|' /C calc'!A0,1\n+HYPERLINK(\"http://x\"),2\n-2+3,3\n@SUM(A1),4\n\t=1+1,5\n"
    text = _text(run_layer2(raw, "csv"))
    for line in text.splitlines()[1:]:
        assert not line.lstrip('"').startswith(("=", "+", "-", "@", "\t"))
    assert "cmd|' /C calc'!A0" in text
    assert "SUM(A1)" in text and "@" not in text


def test_csv_utf8_fallback_latin1():
    raw = "village,value\nFaleāsiu,1\nCafé,2\n".encode("utf-8")
    assert "Faleāsiu" in _text(run_layer2(raw, "csv"))
    latin1 = "village,value\nCafé,2\n".encode("latin-1")  # 0xE9 is invalid UTF-8
    text = _text(run_layer2(latin1, "csv"))
    assert "Café" in text
    text.encode("utf-8")  # re-encoded as UTF-8


def test_csv_empty_after_strip_rejected():
    result = run_layer2(b"=\n@\n+-\n", "csv")
    assert not result.passed
    assert result.rejection_reason == "sanitised content is empty — nothing extractable"


# ── 3. JSON ───────────────────────────────────────────────────────────────────

def test_json_valid_passes():
    raw = json.dumps({"country": "WSM", "values": [1, 2.5, None]}).encode()
    result = run_layer2(raw, "json")
    assert result.passed
    assert json.loads(result.sanitised_content) == {"country": "WSM", "values": [1, 2.5, None]}


def test_json_invalid_rejected():
    for raw in (b"{not json", b'{"a": 1} <script>alert(1)</script>', b'{"a": NaN}', b"{}"):
        result = run_layer2(raw, "json")
        assert not result.passed, raw
        assert result.rejection_reason


def test_json_reserialised_output():
    raw = '﻿{ "b" : 2,\n  "a" : "Āpia" }'.encode("utf-8")
    text = _text(run_layer2(raw, "json"))
    # Canonical: sorted keys, no whitespace, UTF-8, BOM gone.
    assert text == '{"a":"Āpia","b":2}'


# ── 4. TXT ────────────────────────────────────────────────────────────────────

def test_txt_clean_passes():
    raw = "Samoa annual catch report\nTotal: 1,204 t\n".encode("utf-8")
    assert _text(run_layer2(raw, "txt")) == "Samoa annual catch report\nTotal: 1,204 t\n"


def test_txt_null_bytes_stripped():
    result = run_layer2(b"ab\x00c\x00\nd", "txt")
    assert result.sanitised_content == b"abc\nd"
    assert run_layer2(b"\x00\x00 \x00", "txt").passed is False


# ── 5. XML ────────────────────────────────────────────────────────────────────

def test_xml_valid_passes():
    raw = (
        b'<?xml version="1.0"?><?xml-stylesheet href="evil.xsl"?>'
        b'<data><!-- hidden instruction --><obs country="WSM" year="2024">1.5</obs></data>'
    )
    text = _text(run_layer2(raw, "xml"))
    assert text == '<data><obs country="WSM" year="2024">1.5</obs></data>'


def test_xml_xxe_attack_rejected():
    attacks = (
        b'<?xml version="1.0"?><!DOCTYPE d [<!ENTITY xxe SYSTEM "file:///etc/passwd">]><d>&xxe;</d>',
        # Billion-laughs style entity expansion.
        b'<?xml version="1.0"?><!DOCTYPE d [<!ENTITY a "aaaa"><!ENTITY b "&a;&a;&a;&a;">]><d>&b;</d>',
    )
    for raw in attacks:
        result = run_layer2(raw, "xml")
        assert not result.passed
        assert "defusedxml" in result.rejection_reason


def test_xml_malformed_rejected():
    result = run_layer2(b"<data><obs>1</data>", "xml")
    assert not result.passed
    assert "malformed XML" in result.rejection_reason


# ── 6. XLSX ───────────────────────────────────────────────────────────────────

def test_xlsx_clean_passes():
    text = _text(run_layer2(_xlsx([["year", "country", "catch_t"], [2024, "WSM", -3.5]]), "xlsx"))
    assert text == "--- SHEET: Catch ---\nyear,country,catch_t\n2024,WSM,-3.5\n"


def test_xlsx_formulae_stripped():
    raw = _xlsx([["a", "b", "total"], [1, 2, "=SUM(A2:B2)"], ["@cmd", "=HYPERLINK(\"x\")", 3]])
    text = _text(run_layer2(raw, "xlsx"))
    # data_only=True: a formula cell yields its cached value (none here), never its source.
    assert "SUM" not in text and "HYPERLINK" not in text and "=" not in text
    assert "@" not in text and "cmd" in text


def test_xlsx_output_is_text_not_xlsx():
    result = run_layer2(_xlsx([["year"], [2024]]), "xlsx")
    assert result.passed
    assert not result.sanitised_content.startswith(b"PK")
    with pytest.raises(zipfile.BadZipFile):
        zipfile.ZipFile(io.BytesIO(result.sanitised_content))
    result.sanitised_content.decode("utf-8")


# ── 7. DOCX ───────────────────────────────────────────────────────────────────

def test_docx_clean_passes():
    text = _text(run_layer2(_docx(["Samoa Fisheries Report", "Total catch 1,204 t"]), "docx"))
    assert text == "Samoa Fisheries Report\nTotal catch 1,204 t"


def test_docx_output_is_text_not_docx():
    result = run_layer2(_docx(["Report body"]), "docx")
    assert result.passed
    assert not result.sanitised_content.startswith(b"PK")
    assert b"<w:" not in result.sanitised_content  # no WordprocessingML survives


def test_docx_empty_document_rejected():
    result = run_layer2(_docx([]), "docx")
    assert not result.passed
    assert result.rejection_reason == "sanitised content is empty — nothing extractable"


# ── 8. PDF ────────────────────────────────────────────────────────────────────

def test_pdf_clean_passes():
    text = _text(run_layer2(_pdf(["Tuna catch 2024", "Skipjack 1204 t"]), "pdf"))
    assert "Tuna catch 2024" in text
    assert "\n\n--- PAGE 2 ---\n\n" in text
    assert text.index("Tuna catch 2024") < text.index("--- PAGE 2 ---") < text.index("Skipjack 1204 t")


def test_pdf_encrypted_rejected():
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(_pdf(["secret"]))))
    writer.encrypt("pw")
    buf = io.BytesIO()
    writer.write(buf)
    result = run_layer2(buf.getvalue(), "pdf")
    assert not result.passed
    assert result.rejection_reason == "encrypted PDF"


def test_pdf_output_is_text_not_pdf():
    result = run_layer2(_pdf(["Ocean temperature"]), "pdf")
    assert result.passed
    out = result.sanitised_content
    assert not out.startswith(b"%PDF")
    # The catalog's JavaScript OpenAction never reaches Layer 3.
    assert b"JavaScript" not in out and b"app.alert" not in out


# ── 9. content_hash ───────────────────────────────────────────────────────────

def test_content_hash_is_sha256_of_sanitised_not_raw():
    raw = b"name,value\r\n=evil(),1\r\n"
    result = run_layer2(raw, "csv")
    assert result.passed
    assert result.content_hash == hashlib.sha256(result.sanitised_content).hexdigest()
    assert result.content_hash != hashlib.sha256(raw).hexdigest()
    assert len(result.content_hash) == 64 and int(result.content_hash, 16) >= 0


def test_identical_content_produces_identical_hash():
    a = run_layer2(b'{"b":1,"a":2}', "json")
    b = run_layer2(b'{ "a": 2, "b": 1 }\n', "json")
    assert a.passed and b.passed
    assert a.content_hash == b.content_hash
    assert run_layer2(b'{"a":3}', "json").content_hash != a.content_hash


# ── 10. Unknown extension and contract ────────────────────────────────────────

def test_unknown_extension_rejected():
    for ext in ("exe", "html", "xls", "docm", "", "../csv"):
        result = run_layer2(b"year,value\n2024,1\n", ext)
        assert result == Layer2Result(
            passed=False, content_hash=None, sanitised_content=None,
            rejection_reason=f"unsupported file extension: {ext}",
        )


CORRUPT_INPUTS = [
    (b"PK\x03\x04garbage", "xlsx"),
    (b"PK\x03\x04garbage", "docx"),
    (b"%PDF-1.4\n%%EOF", "pdf"),
    (b"\xff\xfe\x00garbage", "pdf"),
    (b"not a zip at all", "xlsx"),
    (b"[" * 100_000 + b"]" * 100_000, "json"),  # deep nesting
    (b"<" * 1000, "xml"),
    (_xlsx([["x"]])[:200], "xlsx"),  # truncated archive
]


def test_never_raises_on_corrupt_input():
    for raw, ext in CORRUPT_INPUTS:
        result = run_layer2(raw, ext)
        assert isinstance(result, Layer2Result)
        assert not result.passed, ext


def test_passed_false_always_has_rejection_reason(monkeypatch):
    def _boom(raw: bytes) -> str:
        raise RuntimeError("parser crashed")

    monkeypatch.setitem(layer2._SANITISERS, "txt", _boom)
    cases = [(raw, ext) for raw, ext in CORRUPT_INPUTS] + [
        (b"", "csv"), (b"   \n", "txt"), (b"x", "exe"), (b"text", "txt"), ("str", "txt"),
    ]
    for raw, ext in cases:
        result = run_layer2(raw, ext)  # type: ignore[arg-type]
        assert not result.passed
        assert isinstance(result.rejection_reason, str) and result.rejection_reason.strip()
        assert result.content_hash is None and result.sanitised_content is None
